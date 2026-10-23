"""One song per editable PPTX slide, plus matching PDF pages, at iPad 3:4 ratio."""
import io
import re
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor

W,H=540.0,720.0  # 7.5 x 10 inches
X=22.0
BW=W-2*X
LX=81.0
TX=X+LX+8.5
TW=W-X-9-TX
START,END=108.0,696.0
INK='#203342'
BORDER='#344e64'
WHITE='#ffffff'
COLORS={'verse':'#eaf2df','chorus':'#ffecdd','prechorus':'#f1e5fc',
        'bridge':'#eee5fa','intro':'#e4edf5','solo':'#c31620',
        'instrumental':'#c31620','break':'#f7e6a2','note':'#fff2a8',
        'outro':'#dfeaf6','other':'#edf0f3'}
DARK={'solo','instrumental'}
PDF_REG,PDF_BOLD='Helvetica','Helvetica-Bold'

def fonts():
    global PDF_REG,PDF_BOLD
    pairs=[('C:/Windows/Fonts/arial.ttf','C:/Windows/Fonts/arialbd.ttf'),
           ('/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf','/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf'),
           ('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf','/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'),
           ('/Library/Fonts/Arial.ttf','/Library/Fonts/Arial Bold.ttf')]
    for normal,bold in pairs:
        if Path(normal).exists() and Path(bold).exists():
            pdfmetrics.registerFont(TTFont('ChartRegular',normal))
            pdfmetrics.registerFont(TTFont('ChartBold',bold))
            PDF_REG,PDF_BOLD='ChartRegular','ChartBold'
            return
fonts()

class ChartOverflow(Exception): pass

def kind_of(name):
    s=re.sub('[^a-z]','',name.lower())
    if s.startswith('pre'): return 'prechorus'
    if s.startswith('verse') or re.match(r'^v\d',name.lower()): return 'verse'
    if s.startswith('chorus') or s in ('ch','c'): return 'chorus'
    if s.startswith('bridge'): return 'bridge'
    if 'solo' in s: return 'solo'
    if 'instrumental' in s: return 'instrumental'
    if 'break' in s or 'stop' in s: return 'break'
    if s.startswith('intro'): return 'intro'
    if s.startswith('outro') or s.startswith('ending') or s=='end': return 'outro'
    if s.startswith('note'): return 'note'
    return 'other'

def short(name,kind):
    m=re.search(r'\d+|\b(?:I|II|III|IV|V)\b',name,re.I)
    n=m.group(0).upper() if m else ''
    return {'verse':'V'+n,'chorus':'CH'+n,'prechorus':'PRE','bridge':'BR',
            'intro':'IN','solo':'SOLO','instrumental':'INST','break':'BREAK',
            'outro':'OUT','note':'NOTE','other':name[:6].upper()}[kind]

CHORDS=re.compile(r'^(?:[A-G](?:#|b)?(?:m|maj|min|dim|aug|sus|add)?\d*(?:/[A-G](?:#|b)?)?|N\.C\.)(?:\s+(?:[A-G](?:#|b)?(?:m|maj|min|dim|aug|sus|add)?\d*(?:/[A-G](?:#|b)?)?|N\.C\.))*$',re.I)

# ChordPro-style [Am] markers are kept as anchors until the page is laid out.
INLINE_CHORD = re.compile(r'\[([^\[\]\n]+)\]')

def inline_parts(line):
    """Return visible lyrics and (character offset, chord) anchors, or None."""
    chunks=[]; anchors=[]; cursor=0; length=0
    for match in INLINE_CHORD.finditer(line):
        chord=match.group(1).strip()
        if not CHORDS.fullmatch(chord):
            continue
        preceding=line[cursor:match.start()]
        chunks.append(preceding)
        length+=len(preceding)
        anchors.append((length,chord))
        cursor=match.end()
    if not anchors:
        return None
    chunks.append(line[cursor:])
    return ''.join(chunks),anchors

def inline_wrapped(text,anchors,size,width):
    """Split lyric text into lines while retaining chord anchor offsets."""
    font=PDF_REG
    words=list(re.finditer(r'\S+',text))
    if not words:
        return [('',[(0,chord) for _,chord in anchors])]
    segments=[]; start=words[0].start(); end=words[0].end()
    for word in words[1:]:
        if pdfmetrics.stringWidth(text[start:word.end()],font,size)<=width:
            end=word.end()
        else:
            segments.append((start,end))
            start,end=word.start(),word.end()
    segments.append((start,end))
    # Break unusually long words too, avoiding any out-of-bounds text.
    fitted=[]
    for start,end in segments:
        while start<end and pdfmetrics.stringWidth(text[start:end],font,size)>width:
            cut=start+1
            while cut<end and pdfmetrics.stringWidth(text[start:cut+1],font,size)<=width:
                cut+=1
            fitted.append((start,cut));start=cut
        if start<end: fitted.append((start,end))
    result=[]
    for i,(start,end) in enumerate(fitted):
        lineanchors=[]
        for pos,chord in anchors:
            # A marker in whitespace before the next wrapped word belongs to that word.
            if start<=pos<end or (i==len(fitted)-1 and pos>=start) or (pos<start and (i==0 or pos>=fitted[i-1][1])):
                lineanchors.append((max(0,pos-start),chord))
        result.append((text[start:end],lineanchors))
    return result

def parse(song):
    sections=[]; current=None; pending=''
    for raw in str(song.get('content','')).replace('\r\n','\n').split('\n'):
        line=raw.strip()
        if not line: continue
        mark=re.fullmatch(r'\[(.+?)\]',line)
        if mark:
            if pending and current: current['rows'].append(('chords',pending,'')); pending=''
            name=mark.group(1).strip()
            current={'name':name,'kind':kind_of(name),'rows':[]}
            sections.append(current)
            continue
        if current is None:
            current={'name':'Intro','kind':'intro','rows':[]}; sections.append(current)
        bold_match = re.match(r'^(?:BOLD|NOTE):\s*(.*)$', line, re.I)
        if bold_match:
            # NOTE: remains a legacy alias; neither prefix is printed.
            if bold_match.group(1):
                current['rows'].append(('bold', bold_match.group(1), ''))
        elif re.match(r'^(?:CUE|GUITAR|KEYS|DRUMS|BASS):',line,re.I):
            current['rows'].append(('note',line,''))
        elif line.startswith('!'):
            current['rows'].append(('note',line[1:].strip(),''))
        elif inline_parts(line) is not None:
            current['rows'].append(('inline',line,''))
        elif ' | ' in line and CHORDS.fullmatch(line.split(' | ',1)[0].strip()):
            chord,lyric=line.split(' | ',1)
            current['rows'].append(('lyric',lyric.strip(),chord.strip()))
        elif CHORDS.fullmatch(line):
            if pending: current['rows'].append(('chords',pending,''))
            pending=line
        else:
            current['rows'].append(('lyric',line,pending));pending=''
    if pending and current: current['rows'].append(('chords',pending,''))
    return sections

def wrap(s,font,size,maxwidth):
    s=str(s).replace('→','->')
    words=s.split(); rows=[]; line=''
    for word in words:
        if pdfmetrics.stringWidth(word,font,size)>maxwidth:
            if line: rows.append(line);line=''
            part=''
            for c in word:
                if part and pdfmetrics.stringWidth(part+c,font,size)>maxwidth:
                    rows.append(part);part=c
                else:part+=c
            line=part
        elif not line or pdfmetrics.stringWidth(line+' '+word,font,size)<=maxwidth:
            line=(line+' '+word).strip()
        else:
            rows.append(line);line=word
    if line:rows.append(line)
    return rows or ['']

def section_layout(sec,y,size):
    kind=sec['kind'];fg=WHITE if kind in DARK else INK
    pad=5*size/11; gap=1.7*size/11; cursor=y+pad; cmds=[]
    for mode,text,chord in sec['rows']:
        if mode=='inline':
            visible,anchors=inline_parts(text)
            chordsize=size*.84
            for lyric,marks in inline_wrapped(visible,anchors,size,TW):
                # Chord positions are measured against the actual rendered lyric font.
                for offset,ch in marks:
                    offset=min(offset,len(lyric))
                    xx=TX+pdfmetrics.stringWidth(lyric[:offset],PDF_REG,size)
                    xx=min(xx,TX+TW-pdfmetrics.stringWidth(ch,PDF_BOLD,chordsize))
                    cmds.append((xx,cursor,ch,chordsize,True,fg))
                cursor+=chordsize*1.22
                cmds.append((TX,cursor,lyric,size,False,fg))
                cursor+=size*1.22
            cursor+=gap
            continue
        if chord:
            s=size*.84
            for l in wrap(chord,PDF_BOLD,s,TW):cmds.append((TX,cursor,l,s,True,fg));cursor+=s*1.1
        s=size*(.93 if mode=='note' else .95 if mode=='chords' else 1.05 if kind in ('break','solo','instrumental') else 1)
        emphasis=mode in ('note','chords','bold') or kind in ('break','solo','instrumental')
        font=PDF_BOLD if emphasis else PDF_REG
        for l in wrap(text,font,s,TW):
            cmds.append((TX,cursor,l,s,emphasis,fg));cursor+=s*(1.22 if mode!='chords' else 1.2)
        cursor+=gap
    if not sec['rows']:
        cmds.append((TX,cursor,'(add song content)',size,False,fg));cursor+=size*1.22
    height=max(27,cursor-y+pad-gap)
    ls=min(10,size*.88)
    labels=wrap(sec['name'].upper(),PDF_BOLD,ls,LX-13)
    ls=min(ls,(height-7)/(1.12*max(len(labels),1)))
    labelcmd=[(X+7,y+7+i*ls*1.12,l,ls,True,fg) for i,l in enumerate(labels)]
    return {'y':y,'h':height,'kind':kind,'label':labelcmd,'cmds':cmds}

def layout(song,index=1,total=1):
    title=(song.get('title') or 'Untitled Song').strip()
    sections=parse(song) or [{'name':'Verse 1','kind':'verse','rows':[('note','Add lyrics, chords and cues in the editor.','')]}]
    for size in (round(11.8 - i * .2, 1) for i in range(15)):
        y=START; boxes=[]
        for sec in sections:
            b=section_layout(sec,y,size);boxes.append(b);y+=b['h']+3.1
        if y-3.1<=END:
            ts=20.2
            while ts>12 and pdfmetrics.stringWidth(title,PDF_BOLD,ts)>440:ts-=.5
            return {'title':title,'artist':song.get('artist',''),'index':index,'total':total,
                    'title_size':ts,'sections':sections,'boxes':boxes,'font':size}
    raise ChartOverflow(f'"{title}" has too much text for a single readable slide. Shorten repeated lyrics or notes, then try again.')

def rgb(h):
    h=h.lstrip('#');return RGBColor(int(h[:2],16),int(h[2:4],16),int(h[4:6],16))

def shape(slide,x,y,w,h,fill,border=None):
    s=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x/72),Inches(y/72),Inches(w/72),Inches(h/72))
    s.fill.solid();s.fill.fore_color.rgb=rgb(fill)
    if border:s.line.color.rgb=rgb(border);s.line.width=Pt(.7)
    else:s.line.fill.background()

def text_shape(slide,x,y,text,size,bold=False,color=INK,width=TW,align=PP_ALIGN.LEFT):
    s=slide.shapes.add_textbox(Inches(x/72),Inches(y/72),Inches(width/72),Inches(size*1.6/72))
    tf=s.text_frame;tf.clear();tf.word_wrap=False
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.TOP
    p=tf.paragraphs[0];p.alignment=align;p.space_before=p.space_after=Pt(0);p.line_spacing=1
    r=p.add_run();r.text=text.replace('→','->');r.font.name='Arial';r.font.size=Pt(size)
    r.font.bold=bold;r.font.color.rgb=rgb(color)

def roadmap(lay):
    sections=lay['sections'];n=len(sections)
    full=BW-5;w=(full-(n-1)*2)/n
    for i,s in enumerate(sections):
        yield X+i*(w+2),w,short(s['name'],s['kind']),s['kind']

def draw_ppt(slide,lay):
    shape(slide,0,0,W,H,WHITE)
    text_shape(slide,X,19,lay['title'],lay['title_size'],True,width=460)
    text_shape(slide,X,48,lay['artist'],10.5,color='#62717d',width=460)
    text_shape(slide,W-77,48,f"{lay['index']} / {lay['total']}",9.3,color='#6b7782',width=55,align=PP_ALIGN.RIGHT)
    for x,w,label,kind in roadmap(lay):
        shape(slide,x,75,w,23,COLORS[kind])
        if w>=21:
            ss=min(9.5,max(5.4,w/max(1,len(label))*1.5))
            text_shape(slide,x+1,80,label,ss,True,WHITE if kind in DARK else INK,max(1,w-2),PP_ALIGN.CENTER)
    for b in lay['boxes']:
        shape(slide,X,b['y'],BW,b['h'],COLORS[b['kind']],BORDER)
        shape(slide,X+LX,b['y'],.75,b['h'],'#324a60')
        for x,y,t,size,bold,color in b['cmds']:
            text_shape(slide,x,y,t,size,bold,color)
        for x,y,t,size,bold,color in b['label']:
            text_shape(slide,x,y,t,size,bold,color,LX-10)

def draw_pdf_text(pdf,x,y,t,size,bold,color):
    pdf.setFont(PDF_BOLD if bold else PDF_REG,size)
    pdf.setFillColor(HexColor(color))
    pdf.drawString(x,H-y-size*.92,t.replace('→','->'))

def draw_pdf(pdf,lay):
    pdf.setFillColor(HexColor(WHITE));pdf.rect(0,0,W,H,stroke=0,fill=1)
    draw_pdf_text(pdf,X,19,lay['title'],lay['title_size'],True,INK)
    draw_pdf_text(pdf,X,48,lay['artist'],10.5,False,'#62717d')
    pdf.setFillColor(HexColor('#6b7782'));pdf.setFont(PDF_REG,9.3)
    pdf.drawRightString(W-X,H-48-9,f"{lay['index']} / {lay['total']}")
    for x,w,label,kind in roadmap(lay):
        pdf.setFillColor(HexColor(COLORS[kind]));pdf.rect(x,H-98,w,23,fill=1,stroke=0)
        if w>=21:
            ss=min(9.5,max(5.4,w/max(1,len(label))*1.5))
            pdf.setFillColor(HexColor(WHITE if kind in DARK else INK))
            pdf.setFont(PDF_BOLD,ss)
            pdf.drawCentredString(x+w/2,H-80-ss*.92,label)
    for b in lay['boxes']:
        pdf.setFillColor(HexColor(COLORS[b['kind']]))
        pdf.setStrokeColor(HexColor(BORDER));pdf.setLineWidth(.7)
        pdf.rect(X,H-b['y']-b['h'],BW,b['h'],fill=1,stroke=1)
        pdf.setFillColor(HexColor('#324a60'))
        pdf.rect(X+LX,H-b['y']-b['h'],.75,b['h'],fill=1,stroke=0)
        for cmd in b['cmds']+b['label']:draw_pdf_text(pdf,*cmd)

def pptx(songs):
    layouts=[layout(s,i+1,len(songs)) for i,s in enumerate(songs)]
    p=Presentation();p.slide_width=Inches(W/72);p.slide_height=Inches(H/72)
    for lay in layouts:draw_ppt(p.slides.add_slide(p.slide_layouts[6]),lay)
    out=io.BytesIO();p.save(out);return out.getvalue()

def pdf(songs):
    layouts=[layout(s,i+1,len(songs)) for i,s in enumerate(songs)]
    out=io.BytesIO();c=canvas.Canvas(out,pagesize=(W,H),pageCompression=1)
    c.setTitle('Live Show Charts')
    for lay in layouts:draw_pdf(c,lay);c.showPage()
    c.save();return out.getvalue()

def preview(song):
    import fitz
    doc=fitz.open(stream=pdf([song]),filetype='pdf')
    result=doc[0].get_pixmap(matrix=fitz.Matrix(1.65,1.65),alpha=False).tobytes('png')
    doc.close();return result
